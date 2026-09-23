from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "研究思路8-25-完善版.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    paragraph.paragraph_format.space_after = Pt(2)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: Iterable[str], rows: Iterable[Iterable[str]], widths: list[float] | None = None) -> None:
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(list(headers)))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(header_cells[i], header, bold=True, color="FFFFFF")
        set_cell_shading(header_cells[i], "1F4E79")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if len(table.rows) % 2 == 0:
                set_cell_shading(cells[i], "EAF2F8")
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph(style="Body Text")
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing = 1.28
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.add_run(text)


def add_number(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.add_run(text)


def add_formula(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Formula")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Python interprets sequences such as ``\frac`` and ``\tilde`` as control
    # characters unless the source literal is raw; restore them before XML write.
    controls = {"\x08": "\\b", "\x0b": "\\v", "\x0c": "\\f", "\t": "\\t"}
    safe_text = "".join(controls.get(char, char) for char in text)
    paragraph.add_run(safe_text)


def add_note(doc: Document, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "FFF2CC")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    p.add_run(f"{title}：").bold = True
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "等线"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    normal.font.size = Pt(10.5)
    body = styles["Body Text"]
    body.font.name = "等线"
    body._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    body.font.size = Pt(10.5)
    for name, size, color in [("Title", 22, "1F4E79"), ("Heading 1", 16, "1F4E79"), ("Heading 2", 13, "2F75B5"), ("Heading 3", 11, "5B9BD5")]:
        style = styles[name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    if "Formula" not in [s.name for s in styles]:
        formula = styles.add_style("Formula", WD_STYLE_TYPE.PARAGRAPH)
    else:
        formula = styles["Formula"]
    formula.font.name = "Cambria Math"
    formula._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    formula.font.size = Pt(10.5)
    formula.font.italic = True

    header = section.header.paragraphs[0]
    header.text = "市场状态感知的多尺度因子表征与动态组合研究方案"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(128, 128, 128)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("研究性方法文档 | 版本：完善版")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(128, 128, 128)


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.space_after = Pt(36)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("市场状态感知的多尺度因子表征\n与动态组合研究方案").bold = True
    p.runs[0].font.name = "微软雅黑"
    p.runs[0]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    p.runs[0].font.size = Pt(24)
    p.runs[0].font.color.rgb = RGBColor(31, 78, 121)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run("基于强化学习生成因子、行为聚类与层次化 MoE 的研究方法论").font.size = Pt(13)
    doc.add_paragraph("\n")
    add_table(doc, ["文档属性", "内容"], [
        ("文档类型", "研究方法论与实验设计方案"),
        ("研究对象", "已有 PPO / 强化学习生成因子及其动态组合"),
        ("研究阶段", "独立研究系统，不依赖前端页面"),
        ("核心方法", "因子编码、行为 K-Means、市场状态矩阵、MACoE、PACoE、层次化 MoE"),
        ("验证原则", "时间顺序切分、严格防止未来数据泄漏、带交易成本样本外回测"),
    ], [4.0, 11.0])
    doc.add_paragraph("\n")
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run("本文件用于研究设计、实验复现和工程实现指导，不构成投资建议。\n")
    p3.add_run("生成日期：2026年9月2日").font.size = Pt(10)
    doc.add_page_break()


def add_toc(doc: Document) -> None:
    add_heading(doc, "目录", 1)
    p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    p._p.append(fld)
    doc.add_paragraph("提示：在 Word 中右键目录并选择“更新域”，即可刷新页码。")
    doc.add_page_break()


def build_document() -> Document:
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_toc(doc)

    add_heading(doc, "摘要", 1)
    add_body(doc, "本研究面向由强化学习等自动化方法生成的候选因子池，提出一套市场状态感知的动态因子组合框架。研究不把单个因子的历史回测排名直接当作最终结论，而是先从多尺度预测能力、稳定性、覆盖率、交易成本和持仓行为等维度构造因子表征，再以因子行为序列进行 K-Means 聚类，获得具有相似表现模式的因子群。在此基础上，使用市场技术与情绪指标构造市场状态矩阵，并通过 MACoE 获取市场环境对各因子的引导分数，通过 PACoE 获取基于近期真实绩效和因子互补性的绩效分数，最后采用层次化 MoE 和动态 Top-k 机制完成因子选择与组合。")
    add_body(doc, "本方案特别强调研究有效性而非模型复杂度。所有滚动绩效特征、聚类参数、路由器参数和组合权重都必须在训练窗口内确定；验证集用于选择超参数和模型版本，测试集只进行一次最终样本外评估。回测中显式计入手续费、滑点、换手和不可交易约束，并通过静态等权、单一类别、随机选择、无状态路由等基准检验复杂模型是否真正带来增益。")
    add_note(doc, "重要边界", "文档中的 MACoE、PACoE 和层次化 MoE 是待实验验证的方法组件，不应在尚未完成训练和样本外检验前表述为已经有效。")

    add_heading(doc, "1 研究背景与问题定义", 1)
    add_heading(doc, "1.1 研究背景", 2)
    add_body(doc, "强化学习因子生成器能够在给定字段、算子和质量约束下搜索大量表达式，但自动生成因子通常存在三个问题：经济含义不稳定、因子之间高度相关、历史表现依赖市场环境。单个因子的全样本平均 IC 或回测收益，可能掩盖其在牛市、熊市、震荡或高波动阶段的差异。因而，本研究将因子视为一组随时间变化的预测工具，而不是固定不变的“好因子”或“坏因子”。")
    add_heading(doc, "1.2 核心研究问题", 2)
    add_number(doc, "如何把可解释性较弱的强化学习因子转换为可比较、可聚类的数值表征？")
    add_number(doc, "如何识别市场所处的状态，并避免使用未来信息定义状态？")
    add_number(doc, "不同因子类别在不同市场环境下是否存在稳定的相对优势？")
    add_number(doc, "能否同时利用市场环境信息和近期绩效信息，动态选择具有互补性的因子？")
    add_number(doc, "动态组合在计入换手和交易成本后，是否仍然优于简单基准？")
    add_heading(doc, "1.3 研究假设", 2)
    add_table(doc, ["编号", "研究假设", "可检验结果"], [
        ("H1", "具有相似行为序列的因子可以形成相对稳定的行为簇。", "簇内相关性高于簇间相关性，且聚类在时间滚动中具有稳定性。"),
        ("H2", "市场状态会改变因子类别的相对表现。", "不同状态下类别条件 Rank IC 或长短收益排序存在差异。"),
        ("H3", "市场引导和绩效引导结合比单一路由更稳健。", "融合模型在验证和测试期的风险调整收益更稳定。"),
        ("H4", "考虑相关性和换手约束可以提升净收益质量。", "成本后 Sharpe、最大回撤或 Calmar 优于无约束组合。"),
    ], [1.2, 6.0, 7.8])

    add_heading(doc, "2 总体研究框架", 1)
    add_body(doc, "整体流程由数据层、表征层、状态层、路由层、组合层和验证层组成。各层通过带 snapshot_id、run_id、配置和输入哈希的研究产物连接，保证一次实验能够被复现和审计。")
    add_table(doc, ["阶段", "输入", "主要处理", "输出"], [
        ("数据准备", "行情、股票池、因子目录、已有指标", "导出独立快照、校验主键、记录来源", "market.parquet、factor_catalog.parquet、manifest"),
        ("逐因子计算", "因子表达式与原始行情", "按表达式计算每个日期和股票的因子值", "factor_values/factor_id=*/values.parquet"),
        ("因子编码", "因子值、未来收益、已有回测结果", "多尺度统计、标准化、缺失处理", "因子编码矩阵与行为序列"),
        ("行为聚类", "因子 IC/Rank IC/斜率序列", "训练期标准化、K-Means、稳定性检验", "cluster_labels、cluster_centers"),
        ("状态识别", "市场技术、广度、情绪指标", "滚动窗口、状态模型或规则", "market_states.parquet"),
        ("动态路由", "市场窗口、因子绩效窗口、聚类标签", "MACoE、PACoE、Group/Factor Router", "类别分数、因子分数、mask、权重"),
        ("组合回测", "因子值、动态权重、行情约束", "下一交易时点执行、费用和滑点", "净值、持仓、交易、绩效指标"),
    ], [2.2, 4.0, 5.3, 3.5])
    add_heading(doc, "2.1 研究对象与符号", 2)
    add_table(doc, ["符号", "含义"], [
        ("t", "交易日索引"),
        ("i", "股票索引"),
        ("k", "候选因子索引，候选池规模为 K"),
        ("r", "预测周期，例如 20 个交易日"),
        ("y_{i,t}^{(r)}", "股票 i 在 t 日之后 r 日的未来收益"),
        ("x_{i,t,k}", "因子 k 在 t 日对股票 i 的因子值"),
        ("s_t", "t 日市场状态"),
        ("g_{t,k}", "MACoE 对因子 k 的市场引导分数"),
        ("p_{t,k}", "PACoE 对因子 k 的绩效引导分数"),
        ("w_{t,k}", "t 日组合中因子 k 的权重"),
    ], [4.0, 11.0])

    add_heading(doc, "3 因子候选池构建", 1)
    add_heading(doc, "3.1 因子来源", 2)
    add_body(doc, "候选因子可以来自 PPO/AlphaGen 等强化学习生成器，也可以来自人工公式、遗传编程或已有研究。进入本研究的因子必须保留表达式、生成运行编号、股票池、预测方向、接受状态和来源签名。研究系统只把原系统作为数据来源，通过导出生成独立快照；后续研究模块不得隐式导入原系统运行时。")
    add_heading(doc, "3.2 候选池准入", 2)
    add_table(doc, ["检查项", "原则", "处理方式"], [
        ("表达式合法性", "字段、算子和参数必须可解析。", "无法解析则剔除并记录错误。"),
        ("覆盖率", "因子值覆盖率达到预设下限。", "低覆盖因子不能直接进入模型。"),
        ("样本数", "每天有效截面至少包含最小股票数。", "不足的日期只记为缺失，不填零。"),
        ("预测方向", "正向和反向因子方向必须统一或显式记录。", "根据训练期方向转换，不能查看测试期决定。"),
        ("极值与异常", "避免无穷、极端数值和常数序列。", "截尾、标准化或剔除，并保留规则。"),
        ("可交易性", "因子值和目标收益必须满足信息时点约束。", "使用 t 日可得值预测 t+1 以后收益。"),
    ], [3.0, 6.0, 6.0])
    add_heading(doc, "3.3 候选池规模控制", 2)
    add_body(doc, "原始因子数量可以较大，但 MACoE/PACoE 不宜直接对上千个高度相关因子进行复杂注意力和动态选择。建议采用两级候选池：先对全部 accepted 因子进行质量过滤和行为聚类，再从每个簇中选取若干代表因子，形成约 20 至 50 个的模型候选池。代表因子的选择只使用训练期指标，并同时考虑绩效和互相关。")
    add_note(doc, "规模建议", "当前系统已有 1,264 个因子逐因子值。第一轮完整模型实验建议压缩至 20-50 个代表因子；原始 1,264 个因子保留用于稳健性分析和后续扩展。")

    add_heading(doc, "4 因子值计算与预处理", 1)
    add_heading(doc, "4.1 因子值定义", 2)
    add_body(doc, "对每个因子表达式 f_k，在每个交易日 t 和股票 i 上计算 x_{i,t,k}=f_k(Z_{i,\leq t})，其中 Z 表示原始行情及已知技术字段。滚动算子只能使用当前日及之前的数据。因子值文件统一采用长表格式：date、code、factor_id、value。")
    add_formula(doc, "x_{i,t,k}=f_k(open, close, high, low, volume, amount, technical_{\leq t})")
    add_heading(doc, "4.2 未来收益标签", 2)
    add_body(doc, "若预测周期为 r，则股票级未来收益定义为当前价格到未来价格的收益。为了模拟 t 日收盘后形成信号、下一交易时点执行的场景，标签和交易执行时点必须在实验配置中明确区分。")
    add_formula(doc, "y_{i,t}^{(r)}=\frac{P_{i,t+r}}{P_{i,t}}-1")
    add_body(doc, "最后 r 个交易日由于缺少未来价格，其标签为空，不能填充为 0。若股票中途停牌，必须使用交易日索引和可交易约束处理，不能把自然日偏移误当成交易日偏移。")
    add_heading(doc, "4.3 截面标准化", 2)
    add_body(doc, "为了消除价格量纲和极端值影响，先在每个日期的可投资截面内做截尾，再进行 Z-score 或 Rank-Gaussian 标准化。标准化参数可以是截面自身的统计量，因为它只使用当日可见数据；如果使用跨时间滚动均值和方差，则必须只使用过去窗口。")
    add_formula(doc, "\tilde{x}_{i,t,k}=\frac{clip(x_{i,t,k},q_{0.01},q_{0.99})-\mu_{t,k}}{\sigma_{t,k}+\epsilon}")

    add_heading(doc, "5 多尺度因子表征与因子编码", 1)
    add_heading(doc, "5.1 编码思想", 2)
    add_body(doc, "因子编码的目标不是简单复制一个总分，而是保留因子在不同预测周期、不同时间窗口和不同评价维度上的行为。对每个因子 k，在训练窗口内构造由生成阶段指标、滚动预测指标和回测行为组成的向量 phi_k。")
    add_formula(doc, "phi_k=[IC_{short},IC_{mid},IC_{long},RIC_{short},RIC_{mid},RIC_{long},ICIR,coverage,turnover,cost,drawdown,holding]\n")
    add_heading(doc, "5.2 指标定义", 2)
    add_table(doc, ["指标", "定义", "研究解释"], [
        ("IC", "因子值与未来收益的 Pearson 相关。", "衡量线性预测能力。"),
        ("Rank IC", "因子排名与未来收益排名的 Spearman 相关。", "衡量排序能力，适合选股场景。"),
        ("ICIR", "日度 IC 均值除以 IC 样本标准差。", "衡量预测能力稳定性。"),
        ("Coverage", "有效因子值占理论样本的比例。", "衡量可用性，不能用 0 代替缺失。"),
        ("回归斜率", "截面回归 y 对标准化因子值的斜率。", "衡量预测幅度和经济量纲。"),
        ("换手率", "相邻调仓期持仓变化规模。", "衡量交易活跃度和成本风险。"),
        ("最大回撤", "净值从历史高点到低点的最大跌幅。", "衡量尾部风险。"),
        ("平均持仓数", "回测期间实际持有股票数量均值。", "衡量组合分散程度。"),
    ], [2.6, 7.0, 5.4])
    add_heading(doc, "5.3 多尺度窗口", 2)
    add_body(doc, "建议至少使用短、中、长三个窗口。例如 20、60、120 个交易日分别描述近期、中期和长期状态。对于每一个窗口，计算 IC 均值、IC 波动、ICIR、Rank IC 均值、Rank ICIR、正 IC 比例、覆盖率和回归斜率等统计量。窗口长度应作为验证集超参数，而不是使用测试集调节。")
    add_heading(doc, "5.4 编码标准化和方向统一", 2)
    add_body(doc, "编码矩阵的每一列在训练期内进行标准化。对于因子方向，如果训练窗口中 Rank IC 均值为负，可将因子值乘以 -1 并记录 direction_flip=true；不能根据测试期方向重新翻转。缺失指标保留 NaN，并在模型输入层采用显式 mask 或训练期统计量填充。")

    add_heading(doc, "6 基于行为序列的 K-Means 因子聚类", 1)
    add_heading(doc, "6.1 聚类输入", 2)
    add_body(doc, "与只按表达式字段分类不同，行为聚类使用因子随时间变化的绩效序列。对每个因子 k，构造按日期排列的行为向量，例如 b_k=[IC_{k,t_1},...,IC_{k,t_n},RankIC_{k,t_1},...,RankIC_{k,t_n}]，也可以使用多尺度窗口摘要和因子收益序列。所有行为特征必须由训练期数据产生。")
    add_formula(doc, "b_k=[z(IC_{k,t}),z(RankIC_{k,t}),z(beta_{k,t}),z(LS_{k,t})]_{t\in T_{train}}")
    add_heading(doc, "6.2 K-Means 目标函数", 2)
    add_formula(doc, "min_{C,mu} \sum_{k=1}^{K} ||b_k-mu_{c_k}||_2^2")
    add_body(doc, "其中 c_k 是因子 k 所属簇，mu_c 是第 c 个簇的中心。实现时需要固定随机种子、初始化方法、最大迭代次数和缺失处理规则。因子行为向量的距离可以使用欧氏距离，也可以在实验中比较相关距离或余弦距离，但必须把距离定义记录在配置中。")
    add_heading(doc, "6.3 簇数 K 的选择", 2)
    add_body(doc, "K 不应根据测试集收益选择。训练期可以使用肘部法、轮廓系数、Calinski-Harabasz 指标和簇内外相关性进行候选比较；验证集用于确认 K 的稳健性。若 K 过小，不同策略行为被强行合并；若 K 过大，类别会失去统计意义且每簇样本太少。建议先比较 K=4、6、8、10。")
    add_heading(doc, "6.4 聚类稳定性", 2)
    add_body(doc, "需要进行重采样稳定性、滚动窗口稳定性和随机种子稳定性检查。可用 Adjusted Rand Index、Normalized Mutual Information 和簇中心相关性比较不同训练窗口产生的标签。如果同一因子在相邻训练窗口频繁跳簇，应将聚类作为描述性标签，而不是直接作为强约束。")
    add_heading(doc, "6.5 簇的解释", 2)
    add_body(doc, "聚类编号本身没有经济含义，不能直接命名为动量簇或反转簇。完成聚类后，根据训练期簇中心在 IC、Rank IC、换手、波动等维度的统计特征进行事后描述，并保留 category_source=behavior_kmeans。若需要经济含义标签，必须由研究者基于表达式和行为共同解释。")
    add_heading(doc, "6.6 代表因子选择", 2)
    add_body(doc, "每个簇内选择若干代表因子，优先使用训练期综合绩效、稳定性、覆盖率和成本后的分数，同时限制代表因子之间的 Spearman 相关。代表因子选择是候选池压缩步骤，不能使用验证或测试期结果。")

    add_heading(doc, "7 市场状态矩阵构建", 1)
    add_heading(doc, "7.1 市场状态的定义", 2)
    add_body(doc, "市场状态不是对未来收益的直接分类，而是对当前可观察环境的压缩表示。状态输入应来自市场或股票池的历史行情、成交、广度和情绪指标。若使用指数或外部宏观数据，必须记录数据发布日期和可用时点。")
    add_heading(doc, "7.2 40 日市场窗口", 2)
    add_body(doc, "按照原研究思路，对每个决策日 t 回看最近 tau=40 个交易日，形成市场状态矩阵 Gamma_t。矩阵的行是历史日期，列是市场特征。原方案中的 72 个指标可以作为目标规模，但在实际实验中应先从独立数据快照中确认字段可用性。")
    add_formula(doc, "Gamma_t=[gamma_{t-39},gamma_{t-38},...,gamma_t]^T \in R^{40\times m}")
    add_table(doc, ["指标组", "示例", "作用"], [
        ("趋势", "市场收益、MA5/MA20/MA60、均线斜率、价格距均线", "描述方向和趋势强度。"),
        ("波动", "收益标准差、ATR、横截面离散度、波动变化率", "描述风险和状态切换。"),
        ("流动性", "成交量均线、成交额变化、换手率分位数", "描述资金活跃度和交易容量。"),
        ("市场广度", "上涨比例、下跌比例、创新高/低比例、涨跌停数量", "描述参与度和内部结构。"),
        ("情绪", "强弱股票比例、连续上涨/下跌家数、极端收益比例", "描述风险偏好和拥挤程度。"),
        ("外部环境", "可获得的指数、利率、商品或宏观变量", "只在发布时间可追溯时使用。"),
    ], [3.0, 7.0, 5.0])
    add_heading(doc, "7.3 滚动标准化", 2)
    add_body(doc, "市场指标应按过去窗口进行标准化，避免用全样本均值和方差把测试期信息泄漏到训练期。对于缺失的外部指标，可以使用最近一次已发布值，但必须标记发布日期和填充规则。")
    add_formula(doc, "\hat{gamma}_{j,t}=\frac{gamma_{j,t}-mean(gamma_{j,t-L:t-1})}{std(gamma_{j,t-L:t-1})+epsilon}")
    add_heading(doc, "7.4 状态识别方法", 2)
    add_body(doc, "第一阶段可以使用可解释规则标签作为基线，例如根据市场趋势和波动阈值识别 bull、bear、sideways、high_volatility。第二阶段可以在训练期拟合 K-Means 或 HMM，验证和测试期只调用已拟合模型。无论采用规则还是聚类，都必须保存 state_method、fit_window_end、特征列、阈值或模型参数。")

    add_heading(doc, "8 MACoE：市场引导注意力中心", 1)
    add_heading(doc, "8.1 目标", 2)
    add_body(doc, "MACoE 用来回答：根据最近市场环境，候选池中的每个因子在当前时点应该有多重要。它关注的是市场叙事和环境结构，不直接等同于近期绩效。输出是长度为 K 的市场引导分数向量。")
    add_heading(doc, "8.2 时间注意力分支", 2)
    add_body(doc, "对 Gamma_t 的 40 个历史时点进行多头注意力，让模型学习近期环境变化和不同市场指标之间的时间依赖。例如最近五日波动突然升高时，模型可以提高这些日期的注意力权重，而不是对 40 日数据简单平均。")
    add_formula(doc, "Attention(Q,K,V)=softmax(QK^T/sqrt(d_k))V")
    add_formula(doc, "gamma_1(t)=MeanPool(MultiHeadAttention(Gamma_t))")
    add_heading(doc, "8.3 链式专家网络 CoE", 2)
    add_body(doc, "将 40 日矩阵沿时间维求平均，得到市场摘要 gamma_input。设置 E=8 个两层 MLP 专家，每个专家使用 Linear-ReLU-Dropout 结构。Router 根据 gamma_input 生成专家权重，上一轮专家输出作为下一轮输入，重复 D=4 轮，获得非线性市场摘要 gamma_2。")
    add_formula(doc, "a^{(d)}=softmax(R^{(d)}(h^{(d-1)})),\quad h^{(d)}=sum_{e=1}^{E}a_e^{(d)}Expert_e^{(d)}(h^{(d-1)})")
    add_heading(doc, "8.4 市场分数输出", 2)
    add_body(doc, "将时间摘要 gamma_1 与专家摘要 gamma_2 拼接并经过融合层，再输入 Scoring MoE，输出 K 个因子分数。该分数在训练期通过下游预测损失间接学习，不应在训练标签之外额外使用测试期表现。")
    add_formula(doc, "g_t=ScoringMoE(Linear([gamma_1(t);gamma_2(t)]))\in R^K")
    add_note(doc, "复杂度控制", "MACoE 的 8 个专家和 4 轮链式结构应作为论文复现配置；第一轮实验可使用更小的专家数和深度，并把规模写入实验编号。")

    add_heading(doc, "9 PACoE：绩效引导注意力中心", 1)
    add_heading(doc, "9.1 目标", 2)
    add_body(doc, "PACoE 用来回答：不考虑市场叙事，哪些因子在最近窗口中预测得更准，并且与其他已选因子具有互补性。PACoE 的输入是因子级近期绩效特征，而不是原始行情。")
    add_heading(doc, "9.2 九维绩效特征", 2)
    add_body(doc, "对每个因子在滚动窗口内计算 IC、Rank IC 和截面回归斜率三类序列；对每一类序列分别计算均值、标准差和均值/标准差，得到 9 维向量。")
    add_formula(doc, "phi_{k,t}=[mean(IC),std(IC),IR(IC),mean(RIC),std(RIC),IR(RIC),mean(beta),std(beta),IR(beta)]")
    add_body(doc, "这里的绩效窗口必须严格滞后于决策日：t 日生成的 phi_{k,t} 只能使用 t 日以前已经能够计算出来的 IC 和收益标签。若未来 20 日收益在 t+20 才实现，则对应 IC 不能在 t 日提前进入 PACoE。实践中需要明确“信号日”和“标签结算日”，或者采用带标签可用延迟的训练样本。")
    add_heading(doc, "9.3 因子关系注意力", 2)
    add_body(doc, "因子间注意力分支学习近期因子表现之间的关联。例如因子 A 和 B 一直同涨同跌、同准同错，注意力可以把邻居信息传递给每个因子，但这不代表二者必然应该同时入选。")
    add_formula(doc, "Phi_1=MultiHeadAttention(Phi,Phi,Phi)")
    add_heading(doc, "9.4 绩效 CoE 与分数", 2)
    add_body(doc, "绩效 CoE 使用 2 个专家、2 轮链式路由，将 9 维成绩单做非线性提炼。融合注意力分支和 CoE 分支后，Scoring MoE 输出长度为 K 的绩效分数 p_t。为了体现互补性，可以在训练损失中加入因子相关性惩罚，或在最终选择器中对已选因子实施相关性门槛。")
    add_formula(doc, "p_t=ScoringMoE(Linear([Phi_1;Phi_2]))\in R^K")

    add_heading(doc, "10 层次化 MoE 动态因子组合", 1)
    add_heading(doc, "10.1 Group-level Router", 2)
    add_body(doc, "Group-level Router 接收市场状态表征或 MACoE 输出，先对行为 K-Means 产生的 G 个因子群进行打分。它可以输出单一最优簇，也可以输出多个簇的 soft 权重。多簇 soft 路由通常比硬选择更平滑，但需要额外的稀疏或熵约束避免所有簇权重接近平均。")
    add_formula(doc, "q_{t,g}=softmax(MLP_{group}(h_t))_g,\quad g=1,...,G")
    add_heading(doc, "10.2 Group Expert", 2)
    add_body(doc, "每个因子群设置一个 Group Expert。Group Expert 在群内处理因子值、绩效向量和因子关系，输出群内因子分数或群内权重。这样可以把“当前哪个类别有优势”和“类别内哪个具体因子更适合”分开建模。")
    add_heading(doc, "10.3 Factor-level Router", 2)
    add_body(doc, "Factor-level Router 在被选中的群内生成因子级分数。若一个因子属于唯一簇，则其最终分数可以由群权重与群内权重相乘得到；如果使用软聚类，则应按因子属于各簇的概率加权。")
    add_formula(doc, "score_{t,k}=sum_{g=1}^{G}q_{t,g}\cdot r_{t,k|g}")
    add_heading(doc, "10.4 MACoE 与 PACoE 融合", 2)
    add_body(doc, "两路分数分别反映环境适配性和近期真实绩效。使用可学习标量 alpha 进行融合，初始值为 0.5，并通过 sigmoid 保证其在 0 到 1 之间。")
    add_formula(doc, "u_{t,k}=alpha_t g_{t,k}+(1-alpha_t)p_{t,k},\quad alpha_t=sigmoid(a_t)")
    add_note(doc, "方向问题", "MACoE 和 PACoE 的分数必须先在训练期完成尺度对齐，例如 LayerNorm 或 z-score，否则某一路数值范围更大时会虚假主导融合。")

    add_heading(doc, "11 Gumbel Top-k 动态因子选择", 1)
    add_heading(doc, "11.1 训练期探索", 2)
    add_body(doc, "为了让模型探索不同因子子集，可以在训练期对分数加入 Gumbel 噪声。温度 tau 从较高值逐步退火到较低值，使训练早期保留探索，后期逐渐接近离散选择。")
    add_formula(doc, "g_k=-log(-log(U_k)),\quad U_k\sim Uniform(0,1)")
    add_formula(doc, "\tilde{u}_{t,k}=u_{t,k}+tau\cdot g_k")
    add_heading(doc, "11.2 Top-k 与稀疏率", 2)
    add_body(doc, "从 K 个因子中选择 k_top 个因子，其比例可设为 60%-100%，原思路的默认值约为 80%。在候选池 K=20 时，k_top 可从 12、16、20 中比较。k_top 不应只以验证期收益决定，还应同时观察换手、因子集中度和成本后收益。")
    add_formula(doc, "m_{t,k}=1[k\in TopK(\tilde{u}_t)]")
    add_heading(doc, "11.3 权重和残差", 2)
    add_body(doc, "被选因子按照 softmax 后的分数加权，未入选因子被 mask。残差项用于保留一个稳定基线，避免路由器在训练初期完全关闭所有原始信息。")
    add_formula(doc, "z_t=sum_k m_{t,k}\cdot softmax(u_t)_k\cdot x_{t,k}+lambda_res\cdot z_t^{base}")
    add_body(doc, "推理阶段不加入随机 Gumbel 噪声，使用确定性 Top-k，或者根据多次采样的入选频率生成稳定权重。所有推理规则必须固定并记录。")

    add_heading(doc, "12 训练目标与约束函数", 1)
    add_heading(doc, "12.1 预测损失", 2)
    add_body(doc, "基本训练目标是利用组合后的 Alpha 预测未来收益。可以使用股票级 MSE，也可以使用截面排序损失或 IC 最大化目标。文档原思路以未来 20 日收益的 MSE 为主，建议保留 MSE 作为基准，再增加排序损失实验。")
    add_formula(doc, "L_{pred}=\frac{1}{N}\sum_{t,i}(\hat{y}_{i,t}-y_{i,t}^{(20)})^2")
    add_heading(doc, "12.2 换手惩罚", 2)
    add_body(doc, "如果每日因子选择变化过快，模型即使预测损失下降，也可能无法交易。可用相邻日期因子权重的 L1 距离近似因子层换手，并在股票组合层再计算实际持仓换手。")
    add_formula(doc, "L_{turn}=\sum_t||w_t-w_{t-1}||_1")
    add_heading(doc, "12.3 相关性和集中度约束", 2)
    add_formula(doc, "L_{corr}=\sum_{k\neq l}w_{t,k}w_{t,l}|rho_{k,l}|")
    add_formula(doc, "L_{entropy}=-\sum_k w_{t,k}log(w_{t,k}+epsilon)")
    add_body(doc, "相关性惩罚抑制重复因子，熵约束用于控制过度集中。实际使用时需要平衡预测能力和组合稳定性，不能把所有因子都强行平均。")
    add_heading(doc, "12.4 总损失", 2)
    add_formula(doc, "L=L_{pred}+lambda_{turn}L_{turn}+lambda_{corr}L_{corr}+lambda_{sparse}L_{sparse}+lambda_{expert}L_{balance}")
    add_body(doc, "其中 L_sparse 控制 Top-k 稀疏性，L_balance 防止专家长期只有一个被使用。所有 lambda 在训练期固定或由验证集选择，测试集不参与调参。")

    add_heading(doc, "13 Alpha 组合与股票组合转换", 1)
    add_heading(doc, "13.1 因子到 Alpha", 2)
    add_body(doc, "动态因子组合先得到每个股票的综合 Alpha 值。为了避免不同因子量纲影响，应先对各因子做截面标准化，再按动态因子权重合成。")
    add_formula(doc, "alpha_{i,t}=\sum_{k=1}^{K}w_{t,k}\tilde{x}_{i,t,k}")
    add_heading(doc, "13.2 Alpha 到股票权重", 2)
    add_body(doc, "基础方案按 Alpha 排名构造多空或多头组合。多头方案可买入前 q 分位股票并等权；多空方案可多头前 q 分位、空头后 q 分位。若当前研究仅支持股票多头，则需要明确基准、现金处理和未交易股票处理。")
    add_formula(doc, "w_{i,t}^{long}=\frac{1}{|L_t|}1[i\in TopQuantile(alpha_t)]")
    add_heading(doc, "13.3 股票级约束", 2)
    add_table(doc, ["约束", "建议规则"], [
        ("单股权重", "不超过预设上限，例如 2% 或按目标持仓数倒数。"),
        ("行业/板块暴露", "如果数据可用，控制相对基准的行业偏离。"),
        ("停牌", "停牌股票不新增仓位，已有仓位按实际可交易状态处理。"),
        ("涨跌停", "无法成交时保留原持仓并记录未成交原因。"),
        ("调仓频率", "按日、周或固定再平衡周期，必须提前设定。"),
        ("现金", "明确现金收益和不可投资资金的处理方式。"),
    ], [5.0, 10.0])

    add_heading(doc, "14 带交易成本约束的样本外回测", 1)
    add_heading(doc, "14.1 执行时点", 2)
    add_body(doc, "回测必须明确决策时点和成交时点。推荐使用 t 日收盘后计算信号，t+1 日开盘或下一可交易时点成交。不能使用 t 日收盘价计算信号后又在 t 日收盘成交，除非明确模拟了可实现的盘中流程。")
    add_heading(doc, "14.2 交易成本模型", 2)
    add_formula(doc, "Cost_t=\sum_i|Delta position_{i,t}|\cdot(c_{fee}+c_{slippage}+c_{impact,i,t})")
    add_body(doc, "最小模型包括手续费和固定滑点；更严格的模型可以让冲击成本随成交额占比、股票流动性和换手率变化。成本参数必须通过配置保存，并在敏感性分析中比较低、中、高三档。")
    add_heading(doc, "14.3 净收益", 2)
    add_formula(doc, "R_t^{net}=R_t^{gross}-Cost_t")
    add_body(doc, "回测同时保存 gross 和 net 两条净值曲线，报告总收益、年化收益、Sharpe、最大回撤、Calmar、日胜率、换手率、总成本、年化成本、交易次数和平均持仓数量。不能只报告收益而隐藏成本。")
    add_heading(doc, "14.4 回测伪代码", 2)
    add_formula(doc, "for t in decision_dates:\n    state = state_model.predict(market_window(t))\n    factor_score = router.predict(state, factor_history(t))\n    selected = deterministic_top_k(factor_score)\n    alpha = combine_factor_values(factor_values[t], selected)\n    target = build_stock_weights(alpha, constraints)\n    executable = apply_suspend_limit_rules(target, holdings[t])\n    cost = turnover_cost(executable, holdings[t], cost_config)\n    holdings[t+1] = executable\n    nav[t+1] = nav[t] * (1 + portfolio_return(executable) - cost)")
    add_note(doc, "关键检查", "回测输出中的每一笔交易都应能追溯到信号日期、状态、入选因子、目标权重、实际成交权重、未成交原因和成本。")

    add_heading(doc, "15 时间切分与防止未来数据泄漏", 1)
    add_heading(doc, "15.1 三段时间切分", 2)
    add_body(doc, "研究采用严格的时间顺序切分。默认可使用 60% 训练、20% 验证、20% 测试，但最终报告应同时记录每段的起止日期和可用样本量。随机切分会把未来市场环境的信息混入训练集，不适合本研究。")
    add_table(doc, ["数据段", "允许用途", "禁止用途"], [
        ("Train", "拟合标准化、K-Means、状态模型、MoE 参数和初始超参数。", "不能反复查看测试期结果。"),
        ("Validation", "选择 K、窗口、专家规模、Top-k、损失权重和成本敏感性方案。", "不能把验证结果宣称为最终样本外表现。"),
        ("Test", "冻结全部配置后运行一次最终评估。", "不能根据测试表现重新选因子、调阈值或修改模型。"),
    ], [3.0, 7.0, 5.0])
    add_heading(doc, "15.2 标签延迟", 2)
    add_body(doc, "未来收益标签存在结算延迟。例如 20 日收益 y_t^{(20)} 只有在 t+20 后才完整可见。因此 PACoE 的近期绩效特征不能把尚未结算的标签当作当前信息。可采用标签成熟日期、purged split 或时间间隔（embargo）处理重叠标签。")
    add_heading(doc, "15.3 运行一致性", 2)
    add_body(doc, "每个运行必须绑定 snapshot_id。因子值运行、状态运行、条件评价运行和组合回测运行的 snapshot_id 不一致时应直接失败。所有输出同时保存配置、输入文件哈希、随机种子、模型版本和代码版本。")

    add_heading(doc, "16 实验方案与对照基线", 1)
    add_heading(doc, "16.1 分级实验", 2)
    add_table(doc, ["实验级别", "模型", "目的"], [
        ("E0", "单因子和全样本等权基线", "确认数据、标签、回测和成本模型正确。"),
        ("E1", "行为 K-Means + 每簇代表因子 + 等权", "验证行为聚类和候选池压缩是否有价值。"),
        ("E2", "状态条件选择 + 固定规则路由", "验证市场状态和类别条件表现。"),
        ("E3", "简化 MACoE + 因子预测", "检验市场环境输入的边际贡献。"),
        ("E4", "简化 PACoE + 因子预测", "检验近期绩效和互补信息的边际贡献。"),
        ("E5", "MACoE + PACoE + 层次化 MoE", "验证完整模型的样本外表现。"),
        ("E6", "完整模型 + Gumbel Top-k + 成本约束", "检验动态稀疏选择和交易成本下的有效性。"),
    ], [2.0, 7.0, 6.0])
    add_heading(doc, "16.2 必要基线", 2)
    add_bullet(doc, "所有候选因子等权组合。")
    add_bullet(doc, "训练期 Rank IC 排名前 N 的静态因子组合。")
    add_bullet(doc, "单一最佳行为簇组合。")
    add_bullet(doc, "随机选择相同数量因子的组合，多次重复报告均值和置信区间。")
    add_bullet(doc, "只有市场状态、没有 PACoE 的路由。")
    add_bullet(doc, "只有近期绩效、没有 MACoE 的路由。")
    add_bullet(doc, "无换手惩罚和有换手惩罚的对照。")
    add_heading(doc, "16.3 统计检验", 2)
    add_body(doc, "对日收益和状态条件 IC 进行 block bootstrap 或按时间区间重采样，报告置信区间。比较两个策略时，可使用配对差收益、均值差、Sharpe 差和回撤差。若进行多模型、多 K、多阈值比较，应记录全部实验，避免只报告最好的一个。")

    add_heading(doc, "17 消融实验与稳健性分析", 1)
    add_table(doc, ["消融项", "去除内容", "需要观察"], [
        ("A1", "去除行为 K-Means，改为表达式结构类别", "聚类是否带来更高簇内一致性。"),
        ("A2", "去除时间注意力", "近期市场变化是否重要。"),
        ("A3", "去除 CoE", "非线性专家结构的边际贡献。"),
        ("A4", "去除 MACoE", "市场环境信息的贡献。"),
        ("A5", "去除 PACoE", "近期绩效信息的贡献。"),
        ("A6", "去除 Gumbel Top-k", "动态稀疏选择是否改善成本后结果。"),
        ("A7", "去除换手、相关性和成本惩罚", "风险约束是否真正有效。"),
        ("A8", "改变训练/验证/测试边界", "结论是否依赖某个特定时期。"),
        ("A9", "改变成本和滑点", "策略对交易摩擦是否敏感。"),
    ], [2.0, 6.0, 7.0])

    add_heading(doc, "18 评价指标体系", 1)
    add_heading(doc, "18.1 因子层指标", 2)
    add_table(doc, ["指标类别", "指标", "判断原则"], [
        ("预测能力", "IC、Rank IC、回归斜率", "关注均值、方向、分布和状态条件差异。"),
        ("稳定性", "ICIR、Rank ICIR、正值比例、滚动稳定性", "不能只看全样本均值。"),
        ("可用性", "Coverage、有效截面样本数", "缺失值保留为空。"),
        ("交易属性", "换手、交易次数、持仓数、成本", "结合净收益判断。"),
        ("相关性", "因子间 Spearman、簇内/簇间相关", "用于去重和互补性分析。"),
    ], [3.0, 6.0, 6.0])
    add_heading(doc, "18.2 组合层指标", 2)
    add_table(doc, ["指标", "公式或说明"], [
        ("总收益", "测试期净值增长率。"),
        ("年化收益", "按交易日年化。"),
        ("Sharpe", "年化平均超额收益 / 年化波动。"),
        ("最大回撤", "净值相对历史峰值的最大跌幅。"),
        ("Calmar", "年化收益 / 最大回撤绝对值。"),
        ("换手率", "相邻持仓变化的规模统计。"),
        ("成本占比", "总成本 / 毛收益或期初资产。"),
        ("状态条件表现", "分别报告 bull、bear、sideways、high_volatility。"),
    ], [4.5, 10.5])
    add_heading(doc, "18.3 选择模型的首要原则", 2)
    add_body(doc, "优先选择测试期表现稳定、成本后仍有收益、不同市场状态下没有灾难性失效的模型，而不是单纯选择测试期总收益最高的模型。对动态路由，必须同时报告入选频率、因子权重集中度、状态切换时的换手和模型置信度。")

    add_heading(doc, "19 数据与工程实现规范", 1)
    add_heading(doc, "19.1 独立数据目录", 2)
    add_formula(doc, "research_data/snapshots/<snapshot_id>/\n  market.parquet\n  universe.csv\n  factor_catalog.parquet\n  factor_metrics.parquet\n  backtest_metrics.parquet\n  manifest.json")
    add_formula(doc, "research_runs/<run_id>/\n  config.json\n  manifest.json\n  market_states.parquet\n  factor_behavior.parquet\n  cluster_labels.parquet\n  router_outputs.parquet\n  portfolio_metrics.parquet")
    add_heading(doc, "19.2 运行元数据", 2)
    add_body(doc, "每次运行至少记录 snapshot_id、输入运行编号、日期边界、预测周期、市场窗口、绩效窗口、K、随机种子、模型超参数、费用配置、软件版本和输出文件哈希。研究结果不能只保存一个最终分数而丢失产生过程。")
    add_heading(doc, "19.3 缺失与异常", 2)
    add_bullet(doc, "缺少必需列或主键重复时立即失败。")
    add_bullet(doc, "无法计算的因子值记录失败原因，不以 0 代替。")
    add_bullet(doc, "截面有效股票数小于 3 时 IC 为空。")
    add_bullet(doc, "状态样本过少时保留指标为空，并在报告中提示统计不稳定。")
    add_bullet(doc, "所有异常处理都写入 manifest 和研究日志。")

    add_heading(doc, "20 端到端研究流程", 1)
    add_heading(doc, "20.1 流程清单", 2)
    steps = [
        "导出原系统行情、股票池、因子目录和已有回测指标，生成不可变 snapshot。",
        "校验日期、股票代码、主键、字段、文件哈希和数据范围。",
        "根据独立快照中的表达式计算逐因子值，保存分区文件。",
        "计算未来收益标签，并严格处理标签成熟日期。",
        "在训练窗口计算多尺度 IC、Rank IC、回归斜率、覆盖率、换手和成本。",
        "对行为序列进行标准化和 K-Means 聚类，选择簇数并保存聚类参数。",
        "从训练期每个簇中选取代表因子，形成 20-50 个模型候选池。",
        "构造最近 40 日市场状态矩阵，训练规则或状态模型。",
        "计算 MACoE 市场引导分数和 PACoE 绩效引导分数。",
        "通过 Group-level Router 选择因子群，再由 Factor-level Router 选择具体因子。",
        "融合两路分数，训练期使用 Gumbel Top-k，推理期使用确定性 Top-k。",
        "将因子组合转换为股票权重，施加单股、可交易性和暴露约束。",
        "按下一交易时点执行，加入手续费、滑点、冲击成本和换手惩罚。",
        "在验证集选择模型和超参数，冻结配置后只在测试集运行一次。",
        "报告因子、类别、状态、组合和成本后的完整指标，并进行消融实验。",
    ]
    for step in steps:
        add_number(doc, step)
    add_heading(doc, "20.2 研究伪代码", 2)
    add_formula(doc, "snapshot = export_and_validate_sources()\nmarket, catalog = load_snapshot(snapshot)\nvalues = materialize_factor_values(market, catalog)\ntarget = make_forward_return(market, horizon=20)\nbehavior = build_rolling_behavior(values, target, train_window)\nclusters = fit_kmeans(behavior.train, K, seed)\nrepresentatives = select_cluster_representatives(behavior.train, clusters)\nstates = fit_market_state_model(market.train, state_config)\nfor t in valid_or_test_dates:\n    market_window = build_market_window(market, t, tau=40)\n    factor_history = build_mature_factor_history(behavior, t)\n    g = macoe(market_window)\n    p = pacoe(factor_history)\n    scores = fuse_scores(g, p)\n    selected = top_k(scores, k_top)\n    stock_weights = alpha_to_stock_weights(values[t], selected)\n    execute_next_bar(stock_weights, costs, constraints)\nreport_metrics_and_audit()")

    add_heading(doc, "21 实际实现阶段划分", 1)
    add_table(doc, ["阶段", "实现内容", "完成标准"], [
        ("P0 数据底座", "快照、校验、逐因子值和未来收益", "数据可复现，主键和哈希校验通过。"),
        ("P1 研究基线", "条件 IC、规则状态、类别评分、等权组合", "形成可解释基线和成本回测。"),
        ("P2 行为聚类", "滚动行为特征、K-Means、代表因子池", "簇稳定性和候选池规模满足要求。"),
        ("P3 动态路由简化版", "规则或小型 MLP 的 Group/Factor Router", "比静态基线有可重复的验证结果。"),
        ("P4 MACoE/PACoE", "注意力和 CoE 模块", "完成单模块消融和训练稳定性检查。"),
        ("P5 完整模型", "融合、Gumbel Top-k、换手和相关性约束", "测试集一次性完成，输出完整审计结果。"),
    ], [3.0, 7.0, 5.0])
    add_body(doc, "工程上不建议直接从 P0 跳到 P5。每个阶段都应该有独立的配置、运行目录、测试和基准结果。只有当 P2 的行为聚类和 P3 的简化路由成立后，才有必要增加论文结构中的多头注意力和链式专家。")

    add_heading(doc, "22 可能遇到的问题与解决方案", 1)
    add_table(doc, ["问题", "原因", "解决方案"], [
        ("训练收益很好，测试崩溃", "模型参数过多、窗口过短或标签泄漏。", "增加时间间隔、减少专家规模、做滚动验证和消融。"),
        ("所有因子高度相关", "生成器重复搜索相似表达式。", "行为聚类、簇内去重、候选池压缩和相关性惩罚。"),
        ("路由每天剧烈变化", "分数噪声大、Top-k 太硬、无换手约束。", "温度退火、权重平滑、换手惩罚、最小持有期。"),
        ("状态类别样本极不均衡", "阈值或聚类导致某状态很少。", "合并稀有状态、使用最小样本阈值并报告不确定性。"),
        ("PACoE 使用了未来标签", "忽略了未来收益的结算延迟。", "标签成熟日期、purged split、embargo 和审计。"),
        ("毛收益高、净收益低", "换手和滑点过高。", "成本敏感训练、降低调仓频率、流动性约束。"),
        ("复杂模型不优于基线", "复杂度没有转化为有效信息。", "保留简单模型作为正式结论，复杂模型只作为研究结果。"),
    ], [4.0, 5.5, 5.5])

    add_heading(doc, "23 研究伦理、局限性与可解释性", 1)
    add_body(doc, "本方法的动态选择能力可能提高模型对历史环境的适应，但也增加了过拟合和解释难度。K-Means 簇是行为相似性描述，不等同于经济因果类别；MACoE/PACoE 分数是模型内部的预测工具，不应被解释为市场状态的唯一真实原因。模型输出仍然依赖数据质量、股票池定义、交易成本假设和预测周期。")
    add_body(doc, "当测试期市场结构发生明显变化时，训练期获得的状态模型和因子行为簇可能失效。因此正式研究报告需要披露测试期范围、市场制度变化、样本覆盖、失败因子比例、不可成交比例和所有重大配置。若结果只在单一股票池或单一预测周期有效，应明确限制外推。")

    add_heading(doc, "24 预期研究产物", 1)
    add_table(doc, ["产物", "内容", "用途"], [
        ("snapshot manifest", "输入文件、来源、字段、行数、哈希", "数据审计和复现。"),
        ("factor behavior", "每个因子的多尺度行为序列和统计量", "编码、聚类和诊断。"),
        ("cluster labels", "因子簇标签、中心、距离和稳定性", "行为类别分析。"),
        ("market states", "日期、状态特征、状态标签和拟合边界", "市场状态路由。"),
        ("router outputs", "类别分数、因子分数、alpha、mask 和权重", "动态选择审计。"),
        ("trades", "信号、目标仓位、实际成交、未成交原因、成本", "回测复核。"),
        ("portfolio metrics", "毛收益、净收益、风险、换手、成本和状态条件指标", "最终实验报告。"),
    ], [4.0, 7.0, 4.0])

    add_heading(doc, "25 总结", 1)
    add_body(doc, "本研究提出的完整逻辑是：先从原始行情真实计算逐因子值，再用多尺度预测指标和交易行为构造因子编码；在训练期用行为序列 K-Means 形成相似因子簇并压缩候选池；以过去 40 个交易日的技术、流动性、广度和情绪指标构造市场状态矩阵；通过 MACoE 获取市场引导分数，通过 PACoE 获取绩效和互补性分数；利用 Group-level Router 先选因子群，再由 Factor-level Router 选择具体因子；训练期以 Gumbel Top-k 探索动态子集，推理期使用确定性选择；最后把 Alpha 转换为股票组合，在下一交易时点执行并计入交易成本。")
    add_body(doc, "该框架的研究价值不在于模型名称或网络层数，而在于能否通过严格的时间边界、可审计数据、合理基线、成本后回测和消融实验证明动态组合确实利用了可泛化的信息。实际实施应遵循从基线到聚类、从简化路由到完整 MoE 的渐进顺序，任何复杂结构都必须接受样本外验证。")

    add_heading(doc, "附录 A：当前独立研究系统的对应关系", 1)
    add_table(doc, ["研究方法模块", "独立系统对应", "当前状态"], [
        ("独立快照", "research_data/snapshots/<snapshot_id>", "已实现"),
        ("逐因子值", "research_runs/values_*/factor_values", "已实现，1,264 个因子可计算"),
        ("市场状态", "research_system states", "已实现规则版"),
        ("条件评价", "research_system conditional", "已实现 IC、Rank IC、ICIR 和时间切分"),
        ("基础选择", "research_system select", "已实现类别评分、因子排序和等权组合"),
        ("行为 K-Means", "待新增 cluster 模块", "本方案定义，待实现"),
        ("MACoE/PACoE", "待新增模型模块", "本方案定义，待实验实现"),
        ("成本样本外回测", "待新增 portfolio/backtest 模块", "本方案定义，待实现"),
    ], [4.5, 6.0, 4.5])

    add_heading(doc, "附录 B：实验报告最低模板", 1)
    add_bullet(doc, "数据：snapshot_id、股票池、日期范围、因子数量、有效覆盖率。")
    add_bullet(doc, "标签：预测周期、信号时点、执行时点、标签成熟和 embargo 规则。")
    add_bullet(doc, "聚类：行为特征、K、距离、随机种子、簇规模和稳定性。")
    add_bullet(doc, "模型：MACoE/PACoE 结构、专家数量、链深、Top-k、损失权重。")
    add_bullet(doc, "回测：手续费、滑点、冲击成本、停牌和涨跌停处理。")
    add_bullet(doc, "结果：Train/Valid/Test 的因子指标、组合指标、状态条件指标。")
    add_bullet(doc, "稳健性：基线、消融、成本敏感性、滚动窗口和随机种子分析。")
    add_bullet(doc, "限制：失败因子、缺失率、不可交易比例、状态不平衡和结论适用范围。")

    return doc


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)
